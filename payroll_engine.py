import calendar
from datetime import datetime
import os
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from config import EXPORTS_DIR
from database import db


class PayrollEngine:
    """محرك احتساب الرواتب مع دعم أيام العمل المخصصة والعطلات والإجازات الطارئة."""

    @staticmethod
    def calculate_employee_month(
        emp: Dict[str, Any], year: int, month: int
    ) -> Dict[str, Any]:
        """
        احتساب أجر الموظف بدقة:
        - أيام العمل الشهرية المقررة: required_work_days (الافتراضي 30 أو حسب العقد)
        - الأجر اليومي = الراتب الأساسي ÷ أيام العمل المقررة
        - فحص أيام العطلات الطارئة والاستثنائية المعتمدة (مدفوعة الأجر)
        - احتساب الامتثال اليومي وفق الورديات
        """
        base_salary = float(emp["base_salary"])
        req_days = int(emp.get("required_work_days") or 30)
        if req_days <= 0:
            req_days = 30

        daily_rate = round(base_salary / req_days, 2)
        shift_type = emp.get("shift_type", "single")
        emp_id = emp["id"]

        num_days = calendar.monthrange(year, month)[1]
        start_date = f"{year:04d}-{month:02d}-01"
        end_date = f"{year:04d}-{month:02d}-{num_days:02d}"

        # استعلام حركات الحضور في هذا الشهر
        logs = db.get_attendance_logs(
            date_from=start_date, date_to=end_date, employee_id=emp_id, limit=5000
        )
        logs_by_date: Dict[str, List[Dict[str, Any]]] = {}
        for log in logs:
            d_str = str(log["date"])
            logs_by_date.setdefault(d_str, []).append(log)

        total_day_credits = 0.0
        daily_details: List[Dict[str, Any]] = []

        # تقييم أيام الشهر (حتى سقف أيام العمل المقررة أو أيام الشهر الفعلية)
        eval_days = min(req_days, num_days)
        for day in range(1, eval_days + 1):
            cur_date_str = f"{year:04d}-{month:02d}-{day:02d}"
            day_logs = logs_by_date.get(cur_date_str, [])
            punch_count = len(day_logs)

            credit = 0.0
            status_text = "غياب كامل"

            # 1. فحص وجود عطلة أو إجازة طارئة معتمدة لهذا اليوم
            emergency = db.is_emergency_day_off(cur_date_str, employee_id=emp_id)
            if emergency:
                if emergency.get("is_paid", 1) == 1:
                    credit = 1.0
                    status_text = f"إجازة طارئة معتمدة ({emergency['reason']})"
                else:
                    credit = 0.0
                    status_text = f"عطلة طارئة غير مدفوعة ({emergency['reason']})"
            else:
                # 2. احتساب الحركات وفق نظام الورديات
                if shift_type == "single":
                    if punch_count >= 2:
                        credit = 1.0
                        status_text = "دوام كامل (مكتمل)"
                    elif punch_count == 1:
                        credit = 0.5
                        status_text = "نصف يوم (نقص بصمة الخروج)"
                    else:
                        credit = 0.0
                        status_text = "غياب تام"
                else:  # Double shift
                    if punch_count >= 4:
                        credit = 1.0
                        status_text = "دوام كامل (ورديتان)"
                    elif punch_count >= 2:
                        credit = 0.5
                        status_text = "نصف يوم (وردية واحدة)"
                    elif punch_count == 1:
                        credit = 0.25
                        status_text = "دوام جزئي (وردية غير مكتملة)"
                    else:
                        credit = 0.0
                        status_text = "غياب تام"

            total_day_credits += credit

            p_map = {
                "IN_1": "دخول 1",
                "OUT_1": "خروج 1",
                "IN_2": "دخول 2",
                "OUT_2": "خروج 2",
            }
            punch_times = ", ".join([f"{p_map.get(l['punch_type'], l['punch_type'])} ({l['time']})" for l in day_logs])
            daily_details.append({
                "date": cur_date_str,
                "punch_count": punch_count,
                "credit": credit,
                "status": status_text,
                "punches": punch_times or ("إجازة طارئة" if emergency else "لا يوجد"),
            })

        days_worked = round(total_day_credits, 2)
        days_absent = round(max(0.0, float(req_days) - days_worked), 2)
        total_deductions = round(days_absent * daily_rate, 2)
        net_pay = round(max(0.0, base_salary - total_deductions), 2)

        return {
            "employee_id": emp_id,
            "fingerprint_id": emp["fingerprint_id"],
            "full_name": emp["full_name"],
            "shift_type": shift_type,
            "base_salary": base_salary,
            "required_work_days": req_days,
            "daily_salary": daily_rate,
            "days_worked": days_worked,
            "days_absent": days_absent,
            "total_deductions": total_deductions,
            "net_pay": net_pay,
            "daily_details": daily_details,
        }

    @classmethod
    def calculate_monthly_payroll(
        cls, year: int, month: int, employee_id: Optional[int] = None
    ) -> List[Dict[str, Any]]:
        if employee_id:
            emp = db.get_employee_by_id(employee_id)
            employees = [emp] if emp else []
        else:
            employees = db.get_all_employees(active_only=True)

        results = []
        for emp in employees:
            record = cls.calculate_employee_month(emp, year, month)
            results.append(record)

        return results

    @classmethod
    def export_payroll_to_excel(
        cls,
        payroll_data: List[Dict[str, Any]],
        year: int,
        month: int,
        company_name: Optional[str] = None,
        custom_filepath: Optional[str] = None,
    ) -> str:
        c_name = company_name or db.get_setting("company_name", "نظام إدارة الموارد البشرية والحضور الذكي")
        month_label = f"{year:04d}-{month:02d}"

        if not custom_filepath:
            filename = f"مسير_رواتب_{month_label}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            filepath = os.path.join(EXPORTS_DIR, filename)
        else:
            filepath = custom_filepath

        wb = openpyxl.Workbook()

        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        total_fill = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")
        
        font_title = Font(name="Tahoma", size=16, bold=True, color="1F4E79")
        font_sub = Font(name="Tahoma", size=11, color="595959")
        font_header = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
        font_body = Font(name="Tahoma", size=10)
        font_total = Font(name="Tahoma", size=11, bold=True)
        
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )
        total_border = Border(
            top=Side(style="thin", color="000000"),
            bottom=Side(style="double", color="000000"),
        )
        align_center = Alignment(horizontal="center", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        # ---------------- ورقة 1: كشف مسير الرواتب ----------------
        ws_summary = wb.active
        ws_summary.title = "مسير الرواتب الشهري"
        ws_summary.views.sheetView[0].showGridLines = True
        ws_summary.views.sheetView[0].rightToLeft = True

        ws_summary["A1"] = c_name
        ws_summary["A1"].font = font_title
        ws_summary["A2"] = f"كشف مسير الرواتب والأجور الشهرية — الفترة: {month_label}"
        ws_summary["A2"].font = font_sub
        ws_summary["A3"] = f"تاريخ الإصدار: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws_summary["A3"].font = font_sub

        start_row = 5
        arabic_headers = [
            "معرف البصمة",
            "اسم الموظف",
            "نوع الدوام",
            "الراتب الأساسي",
            "أيام العمل المقررة",
            "الأجر اليومي",
            "أيام العمل المحتسبة",
            "أيام الغياب",
            "إجمالي الخصومات",
            "صافي الراتب المستحق",
        ]

        for col_idx, header in enumerate(arabic_headers, 1):
            cell = ws_summary.cell(row=start_row, column=col_idx, value=header)
            cell.font = font_header
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = thin_border
            ws_summary.row_dimensions[start_row].height = 28

        current_row = start_row + 1
        total_monthly_sal = 0.0
        total_deduct = 0.0
        total_net = 0.0

        for item in payroll_data:
            ws_summary.cell(row=current_row, column=1, value=item["fingerprint_id"]).alignment = align_center
            ws_summary.cell(row=current_row, column=2, value=item["full_name"]).alignment = align_right

            shift_ar = "مفرد" if item["shift_type"] == "single" else "مزدوج"
            ws_summary.cell(row=current_row, column=3, value=shift_ar).alignment = align_center

            c_base = ws_summary.cell(row=current_row, column=4, value=item["base_salary"])
            c_base.number_format = "#,##0.00"
            c_base.alignment = align_center

            ws_summary.cell(row=current_row, column=5, value=f"{item['required_work_days']} يوم").alignment = align_center

            c_daily = ws_summary.cell(row=current_row, column=6, value=item["daily_salary"])
            c_daily.number_format = "#,##0.00"
            c_daily.alignment = align_center

            ws_summary.cell(row=current_row, column=7, value=item["days_worked"]).alignment = align_center
            ws_summary.cell(row=current_row, column=8, value=item["days_absent"]).alignment = align_center

            c_ded = ws_summary.cell(row=current_row, column=9, value=item["total_deductions"])
            c_ded.number_format = "#,##0.00"
            c_ded.alignment = align_center

            c_net = ws_summary.cell(row=current_row, column=10, value=item["net_pay"])
            c_net.number_format = "#,##0.00"
            c_net.alignment = align_center

            for col_idx in range(1, 11):
                c = ws_summary.cell(row=current_row, column=col_idx)
                c.font = font_body
                c.border = thin_border

            total_monthly_sal += item["base_salary"]
            total_deduct += item["total_deductions"]
            total_net += item["net_pay"]

            ws_summary.row_dimensions[current_row].height = 22
            current_row += 1

        # صف الإجمالي العام
        ws_summary.cell(row=current_row, column=1, value="")
        ws_summary.cell(row=current_row, column=2, value="الإجمالي العام").font = font_total
        ws_summary.cell(row=current_row, column=3, value="")
        
        c_tot_sal = ws_summary.cell(row=current_row, column=4, value=total_monthly_sal)
        c_tot_sal.number_format = "#,##0.00"
        c_tot_sal.font = font_total
        c_tot_sal.alignment = align_center

        ws_summary.cell(row=current_row, column=5, value="")
        ws_summary.cell(row=current_row, column=6, value="")
        ws_summary.cell(row=current_row, column=7, value="")
        ws_summary.cell(row=current_row, column=8, value="")

        c_tot_ded = ws_summary.cell(row=current_row, column=9, value=total_deduct)
        c_tot_ded.number_format = "#,##0.00"
        c_tot_ded.font = font_total
        c_tot_ded.alignment = align_center

        c_tot_net = ws_summary.cell(row=current_row, column=10, value=total_net)
        c_tot_net.number_format = "#,##0.00"
        c_tot_net.font = font_total
        c_tot_net.alignment = align_center

        for col_idx in range(1, 11):
            c = ws_summary.cell(row=current_row, column=col_idx)
            c.fill = total_fill
            c.border = total_border

        ws_summary.row_dimensions[current_row].height = 26

        for col in ws_summary.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if cell.row > 4:
                    max_len = max(max_len, len(val))
            ws_summary.column_dimensions[col_letter].width = max(max_len + 6, 16)

        # ---------------- ورقة 2: تفاصيل الحضور اليومي ----------------
        ws_detail = wb.create_sheet(title="تفاصيل الحضور اليومي")
        ws_detail.views.sheetView[0].showGridLines = True
        ws_detail.views.sheetView[0].rightToLeft = True

        ws_detail["A1"] = f"{c_name} — سجل الحركات والورديات اليومية"
        ws_detail["A1"].font = font_title
        ws_detail["A2"] = f"الفترة: {month_label}"
        ws_detail["A2"].font = font_sub

        detail_headers = [
            "التاريخ",
            "معرف البصمة",
            "اسم الموظف",
            "نظام الدوام",
            "حالة الدوام اليومي",
            "الأجر المحتسب",
            "سجل الحركات المسجلة",
        ]

        detail_start_row = 4
        for col_idx, header in enumerate(detail_headers, 1):
            cell = ws_detail.cell(row=detail_start_row, column=col_idx, value=header)
            cell.font = font_header
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = thin_border
            ws_detail.row_dimensions[detail_start_row].height = 26

        d_row = detail_start_row + 1
        for item in payroll_data:
            shift_ar = "مفرد" if item["shift_type"] == "single" else "مزدوج"
            for d in item.get("daily_details", []):
                ws_detail.cell(row=d_row, column=1, value=d["date"]).alignment = align_center
                ws_detail.cell(row=d_row, column=2, value=item["fingerprint_id"]).alignment = align_center
                ws_detail.cell(row=d_row, column=3, value=item["full_name"]).alignment = align_right
                ws_detail.cell(row=d_row, column=4, value=shift_ar).alignment = align_center
                ws_detail.cell(row=d_row, column=5, value=d["status"]).alignment = align_center
                ws_detail.cell(row=d_row, column=6, value=f"{d['credit']} يوم").alignment = align_center
                ws_detail.cell(row=d_row, column=7, value=d["punches"]).alignment = align_right

                for col_idx in range(1, 8):
                    c = ws_detail.cell(row=d_row, column=col_idx)
                    c.font = font_body
                    c.border = thin_border
                d_row += 1

        for col in ws_detail.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if cell.row > 3:
                    max_len = max(max_len, len(val))
            ws_detail.column_dimensions[col_letter].width = max(max_len + 6, 16)

        wb.save(filepath)
        return filepath


payroll_engine = PayrollEngine()
